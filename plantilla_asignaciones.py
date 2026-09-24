# -- PLANTILLA DE IMPRESIÓN DE ETIQUETAS (HOJA MEMBRETADA V&C) -- #
# Replica el formato de "TJX028 CREMA FACIAL 50 ml.docx": hoja carta con el
# membrete de fondo, un título con la asignación, el subtítulo y el tipo, y
# la etiqueta dentro de un recuadro centrado. Si se conoce la altura del
# contenido, se indica a la izquierda del recuadro, a la altura del contenido.
#
# Todas las medidas están en puntos tipográficos (1 pt = 1/72 in). El mismo
# armado (`calcular_plantilla`) lo usan el PDF y el Word, así ambos coinciden.
import io
import os
from copy import deepcopy
from functools import lru_cache

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

MEMBRETE_PATH = os.path.join("img", "Membrete.jpg")
DPI_MEMBRETE = 200

SUBTITULO = "Muestra de etiqueta con información faltante"
TEXTO_ALTURA = "Altura del contenido"

# Hoja carta y márgenes del Word de referencia.
PAGINA_ANCHO = 612
PAGINA_ALTO = 792
MARGEN_SUPERIOR = 70.85
MARGEN_INFERIOR = 70.85
MARGEN_LADOS = 85.05
# Límite inferior de la etiqueta: arriba de la franja verde del membrete.
LIMITE_INFERIOR = 720

# Encabezado de la hoja (título, subtítulo y tipo).
TAMANO_TITULO = 11
INTERLINEADO_TITULO = 16
TAMANO_SUBTITULO = 14
INTERLINEADO_SUBTITULO = 20
ESPACIO_TRAS_PARRAFO = 10
ESPACIO_ANTES_ETIQUETA = 6

# Recuadro de la etiqueta. En vertical mide 7 × 9 cm como mínimo; en
# horizontal se intercambian las medidas. Crece hacia abajo si el texto no cabe.
ETIQUETA_LADO_CORTO = 198.75
ETIQUETA_LADO_LARGO = 253.85
RELLENO_LADOS = 8
RELLENO_VERTICAL = 14
GROSOR_BORDE = 0.5

# Estilo de cada campo dentro de la etiqueta (fuente Arial).
TAMANO_CAMPO = 10
TAMANO_DESTACADO = 11
TAMANO_CONTENIDO = 14
FACTOR_INTERLINEADO = 1.15
ESPACIO_ENTRE_CAMPOS = 11.5
CAMPOS_DESTACADOS = ("DESCRIPCION", "DENOMINACION", "EAN", "MARCA")
CAMPOS_CONTENIDO = ("CONTENIDO", "CONTENIDO NETO")
# Si el texto no cabe en la hoja se reduce la letra hasta este factor.
ESCALA_MINIMA = 0.6

# Anotación "Altura del contenido" (dos renglones: leyenda y valor).
TAMANO_ANOTACION = 11
INTERLINEADO_ANOTACION = 14

# Fuentes: (nombre en Word, nombre registrado en el PDF, rutas, respaldo PDF).
FUENTES = {
    "campo": ("Arial", "PlantillaArial", ("arial.ttf", "C:/Windows/Fonts/arial.ttf",
              "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"), "Helvetica"),
    "campo_negrita": ("Arial", "PlantillaArialBold", ("arialbd.ttf", "C:/Windows/Fonts/arialbd.ttf",
                      "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"), "Helvetica-Bold"),
    "titulo": ("Calibri", "PlantillaCalibriBold", ("calibrib.ttf", "C:/Windows/Fonts/calibrib.ttf"),
               "Helvetica-Bold"),
    "subtitulo": ("Tw Cen MT", "PlantillaTwCen", ("TCM_____.TTF", "C:/Windows/Fonts/TCM_____.TTF"),
                  "Helvetica"),
}


@lru_cache(maxsize=None)
def _fuente_pdf(clave):
    """Registra la fuente en reportlab (una sola vez) y devuelve su nombre.
    Si no está instalada se usa la fuente de respaldo de reportlab."""
    _, nombre, rutas, respaldo = FUENTES[clave]
    for ruta in rutas:
        try:
            pdfmetrics.registerFont(TTFont(nombre, ruta))
            return nombre
        except Exception:
            continue
    return respaldo


def _envolver(texto, fuente, tamano, ancho_max):
    """Parte el texto en líneas que caben en `ancho_max` puntos, midiendo con
    la fuente real (no por número de caracteres)."""
    lineas = []
    for parrafo in str(texto).splitlines() or [""]:
        actual = ""
        for palabra in parrafo.split():
            prueba = f"{actual} {palabra}" if actual else palabra
            if actual and pdfmetrics.stringWidth(prueba, fuente, tamano) > ancho_max:
                lineas.append(actual)
                actual = palabra
            else:
                actual = prueba
        lineas.append(actual)
    return lineas


def _estilo_campo(campo):
    campo = campo.strip().upper()
    if campo in CAMPOS_CONTENIDO:
        return "campo_negrita", TAMANO_CONTENIDO
    if campo in CAMPOS_DESTACADOS:
        return "campo_negrita", TAMANO_DESTACADO
    return "campo", TAMANO_CAMPO


def _armar_campos(campos_texto, ancho_texto, escala):
    campos = []
    for campo, texto in campos_texto:
        clave, tamano = _estilo_campo(campo)
        tamano *= escala
        campos.append({
            "campo": campo,
            "fuente": clave,
            "tamano": tamano,
            "interlineado": tamano * FACTOR_INTERLINEADO,
            "lineas": _envolver(texto, _fuente_pdf(clave), tamano, ancho_texto),
        })
    return campos


def calcular_plantilla(campos_texto, titulo, tipo=None, altura_contenido=None,
                       orientacion="vertical"):
    """Calcula la posición de todo lo que va en la hoja.

    `campos_texto` es la lista de (campo, texto) en el orden de la norma. El
    resultado lo consumen `guardar_plantilla_pdf` y `guardar_plantilla_docx`."""
    if not campos_texto:
        raise ValueError("No hay campos con datos para generar la etiqueta")

    horizontal = orientacion == "horizontal"
    ancho = ETIQUETA_LADO_LARGO if horizontal else ETIQUETA_LADO_CORTO
    alto_minimo = ETIQUETA_LADO_CORTO if horizontal else ETIQUETA_LADO_LARGO

    encabezado = [{
        "texto": titulo, "fuente": "titulo", "tamano": TAMANO_TITULO,
        "interlineado": INTERLINEADO_TITULO,
    }, {
        "texto": SUBTITULO, "fuente": "subtitulo", "tamano": TAMANO_SUBTITULO,
        "interlineado": INTERLINEADO_SUBTITULO,
    }]
    if tipo:
        encabezado.append({
            "texto": f"Tipo: {tipo}", "fuente": "subtitulo", "tamano": TAMANO_SUBTITULO,
            "interlineado": INTERLINEADO_SUBTITULO,
        })
    y = MARGEN_SUPERIOR
    for parrafo in encabezado:
        parrafo["y"] = y
        y += parrafo["interlineado"] + ESPACIO_TRAS_PARRAFO
    etiqueta_y = y + ESPACIO_ANTES_ETIQUETA - ESPACIO_TRAS_PARRAFO

    # Si el texto es tan largo que el recuadro se saldría de la hoja, se
    # reduce la letra poco a poco en lugar de cortar la etiqueta.
    ancho_texto = ancho - 2 * RELLENO_LADOS
    escala = 1.0
    while True:
        campos = _armar_campos(campos_texto, ancho_texto, escala)
        alto_texto = sum(len(c["lineas"]) * c["interlineado"] for c in campos)
        alto_texto += ESPACIO_ENTRE_CAMPOS * escala * (len(campos) - 1)
        alto = max(alto_minimo, alto_texto + 2 * RELLENO_VERTICAL)
        if etiqueta_y + alto <= LIMITE_INFERIOR or escala <= ESCALA_MINIMA:
            break
        escala = round(escala - 0.05, 2)

    y = RELLENO_VERTICAL
    for i, campo in enumerate(campos):
        if i:
            y += ESPACIO_ENTRE_CAMPOS * escala
        campo["y"] = y  # relativo a la parte superior del recuadro
        y += len(campo["lineas"]) * campo["interlineado"]

    # El renglón del valor queda a la altura de la primera línea del
    # contenido; si la etiqueta no lleva contenido, arriba del recuadro.
    anotacion = None
    if altura_contenido:
        destino = next((c for c in campos if c["campo"].strip().upper() in CAMPOS_CONTENIDO), None)
        y_texto = RELLENO_VERTICAL
        if destino:
            y_texto = max(RELLENO_VERTICAL,
                          destino["y"] + destino["interlineado"] / 2 - 1.5 * INTERLINEADO_ANOTACION)
        anotacion = {"valor": str(altura_contenido), "y": y_texto}

    return {
        "encabezado": encabezado,
        "etiqueta": {
            "x": (PAGINA_ANCHO - ancho) / 2, "y": etiqueta_y,
            "ancho": ancho, "alto": alto, "campos": campos,
        },
        "anotacion": anotacion,
    }


@lru_cache(maxsize=None)
def _membrete_rgb(membrete_path):
    """Devuelve el membrete como JPEG RGB listo para incrustar.

    El archivo original viene en CMYK con un encabezado que python-docx no
    reconoce, y pesa casi 1 MB; se convierte una sola vez en memoria (a
    ~200 dpi, suficiente para imprimir) y se reutiliza en cada etiqueta."""
    if not os.path.exists(membrete_path):
        raise FileNotFoundError(
            f"No se encontró el membrete '{membrete_path}'. Debe estar en la carpeta img "
            "junto a la aplicación."
        )
    with Image.open(membrete_path) as imagen:
        imagen = imagen.convert("RGB")
        ancho_max = int(PAGINA_ANCHO / 72 * DPI_MEMBRETE)
        if imagen.width > ancho_max:
            alto = round(imagen.height * ancho_max / imagen.width)
            imagen = imagen.resize((ancho_max, alto), Image.LANCZOS)
        salida = io.BytesIO()
        imagen.save(salida, format="JPEG", quality=88, dpi=(DPI_MEMBRETE, DPI_MEMBRETE))
    return salida.getvalue()


# ---------------------------------------------------------------- PDF ---- #

def _texto(c, texto, fuente, tamano, x, y_superior, interlineado, centrado=True):
    """Escribe una línea (centrada en `x` o empezando en `x`) dentro de un
    renglón de alto `interlineado`; `y_superior` se mide desde arriba de la hoja."""
    fuente_pdf = _fuente_pdf(fuente)
    ascenso = pdfmetrics.getAscent(fuente_pdf, tamano)
    base = PAGINA_ALTO - (y_superior + (interlineado - tamano * FACTOR_INTERLINEADO) / 2 + ascenso)
    c.setFont(fuente_pdf, tamano)
    if centrado:
        c.drawCentredString(x, base, texto)
    else:
        c.drawString(x, base, texto)


def guardar_plantilla_pdf(plantilla, ruta_salida, membrete_path=MEMBRETE_PATH):
    membrete = _membrete_rgb(membrete_path)
    c = canvas.Canvas(ruta_salida, pagesize=(PAGINA_ANCHO, PAGINA_ALTO))
    c.setTitle(os.path.splitext(os.path.basename(ruta_salida))[0])
    c.drawImage(ImageReader(io.BytesIO(membrete)), 0, 0, PAGINA_ANCHO, PAGINA_ALTO)

    for parrafo in plantilla["encabezado"]:
        _texto(c, parrafo["texto"], parrafo["fuente"], parrafo["tamano"],
                        PAGINA_ANCHO / 2, parrafo["y"], parrafo["interlineado"])

    etiqueta = plantilla["etiqueta"]
    x, y, ancho, alto = etiqueta["x"], etiqueta["y"], etiqueta["ancho"], etiqueta["alto"]
    c.setLineWidth(GROSOR_BORDE)
    c.setFillColorRGB(1, 1, 1)
    c.rect(x, PAGINA_ALTO - y - alto, ancho, alto, stroke=1, fill=1)
    c.setFillColorRGB(0, 0, 0)
    for campo in etiqueta["campos"]:
        for i, linea in enumerate(campo["lineas"]):
            _texto(c, linea, campo["fuente"], campo["tamano"], x + ancho / 2,
                            y + campo["y"] + i * campo["interlineado"], campo["interlineado"])

    anotacion = plantilla["anotacion"]
    if anotacion:
        _texto(c, TEXTO_ALTURA, "titulo", TAMANO_ANOTACION, MARGEN_LADOS,
               y + anotacion["y"], INTERLINEADO_ANOTACION, centrado=False)
        _texto(c, anotacion["valor"], "titulo", TAMANO_ANOTACION, MARGEN_LADOS,
               y + anotacion["y"] + INTERLINEADO_ANOTACION, INTERLINEADO_ANOTACION, centrado=False)

    c.showPage()
    c.save()
    return ruta_salida


# --------------------------------------------------------------- Word ---- #

def _agregar_xml(padre, etiqueta, atributos=None, hijos=()):
    elemento = OxmlElement(etiqueta)
    for clave, valor in (atributos or {}).items():
        elemento.set(qn(clave), valor)
    for hijo_tag, hijo_attrs in hijos:
        _agregar_xml(elemento, hijo_tag, hijo_attrs)
    padre.append(elemento)
    return elemento


def _formatear_parrafo(parrafo, interlineado, antes=0, despues=0, alineacion=WD_ALIGN_PARAGRAPH.CENTER):
    formato = parrafo.paragraph_format
    formato.alignment = alineacion
    formato.space_before = Pt(antes)
    formato.space_after = Pt(despues)
    formato.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    formato.line_spacing = Pt(interlineado)


def _agregar_texto(parrafo, lineas, fuente, tamano):
    nombre_word = FUENTES[fuente][0]
    for i, linea in enumerate(lineas):
        run = parrafo.add_run(linea)
        run.bold = fuente in ("campo_negrita", "titulo")
        run.font.name = nombre_word
        run.font.size = Pt(tamano)
        # Sin esto Word usa la fuente del tema para caracteres acentuados.
        run._r.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), nombre_word)
        if i < len(lineas) - 1:
            run.add_break()


def _emu(puntos):
    return int(Emu(Pt(puntos)))


def _membrete_de_fondo(seccion, membrete):
    """Coloca el membrete en el encabezado como imagen detrás del texto que
    cubre toda la hoja (igual que el Word de referencia)."""
    parrafo = seccion.header.paragraphs[0]
    _formatear_parrafo(parrafo, 1)
    run = parrafo.add_run()
    run.add_picture(io.BytesIO(membrete), width=Pt(PAGINA_ANCHO), height=Pt(PAGINA_ALTO))
    drawing = run._r.find(qn("w:drawing"))
    en_linea = drawing.find(qn("wp:inline"))
    grafico = en_linea.find(qn("a:graphic"))
    doc_pr = en_linea.find(qn("wp:docPr"))

    ancla = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{_emu(PAGINA_ANCHO)}" cy="{_emu(PAGINA_ALTO)}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>'
    )
    ancla.append(deepcopy(doc_pr))
    _agregar_xml(ancla, "wp:cNvGraphicFramePr")
    ancla.append(deepcopy(grafico))
    drawing.replace(en_linea, ancla)


def _celda(celda, ancho, bordes=False):
    celda.width = Pt(ancho)
    tc_pr = celda._tc.get_or_add_tcPr()
    lados = ("top", "left", "bottom", "right")
    borde = {"w:val": "single", "w:sz": "4", "w:space": "0", "w:color": "000000"} if bordes else {"w:val": "nil"}
    _agregar_xml(tc_pr, "w:tcBorders", hijos=[(f"w:{lado}", borde) for lado in lados])
    # Word aplica a toda la fila el margen superior más grande de sus celdas,
    # así que las tres llevan el mismo; lo demás se ajusta con espacio antes.
    margenes = {"top": RELLENO_VERTICAL, "bottom": 0, "left": RELLENO_LADOS if bordes else 0,
                "right": RELLENO_LADOS if bordes else 0}
    _agregar_xml(tc_pr, "w:tcMar", hijos=[
        (f"w:{lado}", {"w:w": str(round(valor * 20)), "w:type": "dxa"}) for lado, valor in margenes.items()
    ])


def guardar_plantilla_docx(plantilla, ruta_salida, membrete_path=MEMBRETE_PATH):
    membrete = _membrete_rgb(membrete_path)
    doc = Document()
    doc.core_properties.title = os.path.splitext(os.path.basename(ruta_salida))[0]
    seccion = doc.sections[0]
    seccion.page_width = Pt(PAGINA_ANCHO)
    seccion.page_height = Pt(PAGINA_ALTO)
    seccion.top_margin = Pt(MARGEN_SUPERIOR)
    seccion.bottom_margin = Pt(MARGEN_INFERIOR)
    seccion.left_margin = seccion.right_margin = Pt(MARGEN_LADOS)
    seccion.header_distance = Pt(35.4)
    seccion.footer_distance = Pt(35.4)
    _membrete_de_fondo(seccion, membrete)

    encabezado = plantilla["encabezado"]
    for i, parrafo in enumerate(encabezado):
        p = doc.add_paragraph()
        despues = ESPACIO_ANTES_ETIQUETA if i == len(encabezado) - 1 else ESPACIO_TRAS_PARRAFO
        _formatear_parrafo(p, parrafo["interlineado"], despues=despues)
        _agregar_texto(p, [parrafo["texto"]], parrafo["fuente"], parrafo["tamano"])

    # Tabla de tres columnas: anotación | etiqueta con borde | vacía. Así el
    # recuadro queda centrado en la hoja y la anotación a su izquierda.
    etiqueta = plantilla["etiqueta"]
    ancho_util = PAGINA_ANCHO - 2 * MARGEN_LADOS
    lateral = (ancho_util - etiqueta["ancho"]) / 2

    tabla = doc.add_table(rows=1, cols=3)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    _agregar_xml(tabla._tbl.tblPr, "w:tblLayout", {"w:type": "fixed"})
    fila = tabla.rows[0]
    # Word suma el margen superior de las celdas a la altura de la fila.
    fila.height = Pt(etiqueta["alto"] - RELLENO_VERTICAL)
    fila.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for columna, ancho in zip(tabla.columns, (lateral, etiqueta["ancho"], lateral)):
        columna.width = Pt(ancho)

    izquierda, centro, derecha = fila.cells
    _celda(izquierda, lateral)
    _celda(centro, etiqueta["ancho"], bordes=True)
    _celda(derecha, lateral)

    fin_anterior = RELLENO_VERTICAL
    for i, campo in enumerate(etiqueta["campos"]):
        p = centro.paragraphs[0] if i == 0 else centro.add_paragraph()
        _formatear_parrafo(p, campo["interlineado"], antes=campo["y"] - fin_anterior)
        _agregar_texto(p, campo["lineas"], campo["fuente"], campo["tamano"])
        fin_anterior = campo["y"] + len(campo["lineas"]) * campo["interlineado"]

    anotacion = plantilla["anotacion"]
    p = izquierda.paragraphs[0]
    if anotacion:
        _formatear_parrafo(p, INTERLINEADO_ANOTACION, antes=anotacion["y"] - RELLENO_VERTICAL,
                           alineacion=WD_ALIGN_PARAGRAPH.LEFT)
        _agregar_texto(p, [TEXTO_ALTURA], "titulo", TAMANO_ANOTACION)
        valor = izquierda.add_paragraph()
        _formatear_parrafo(valor, INTERLINEADO_ANOTACION, alineacion=WD_ALIGN_PARAGRAPH.LEFT)
        _agregar_texto(valor, [anotacion["valor"]], "titulo", TAMANO_ANOTACION)
    else:
        _formatear_parrafo(p, 1)
    _formatear_parrafo(derecha.paragraphs[0], 1)

    # Word exige un párrafo después de la tabla; se deja de 1 pt para que no
    # empuje nada a una segunda hoja.
    _formatear_parrafo(doc.add_paragraph(), 1)

    doc.save(ruta_salida)
    return ruta_salida
